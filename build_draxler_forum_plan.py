from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    from PIL import Image
except ImportError as exc:
    raise SystemExit("Pillow is required for image conversion") from exc


ROOT = Path(r"C:\aero")
ASSETS = ROOT / "draxler_forum_plan_assets"
OUT = ROOT / "Draxler_Forum_Content_Plan.docx"
DERIVED = ROOT / "draxler_forum_plan_derived"
DERIVED.mkdir(exist_ok=True)


def convert_to_png(source_name: str) -> Path:
    source = ASSETS / source_name
    target = DERIVED / f"{source.stem}.png"
    with Image.open(source) as image:
        image.convert("RGB").save(target, "PNG")
    return target


DRX101 = convert_to_png("drx-101.webp")
DRX301 = convert_to_png("drx-301.webp")


def set_run_font(run, size=11, bold=None, italic=None, color=None, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url, color="1F4E79"):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:cs"), "Arial")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_para(doc, text="", style=None, bold_prefix=None, italic=False, space_after=None, keep_with_next=False):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        lead = p.add_run(bold_prefix)
        set_run_font(lead, 11, bold=True)
        remainder = p.add_run(text[len(bold_prefix):])
        set_run_font(remainder, 11, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, 11, italic=italic)
    pf = p.paragraph_format
    pf.line_spacing = 1.14
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if keep_with_next:
        pf.keep_with_next = True
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, 15 if level == 1 else 12.5, bold=True, color=(0, 0, 0))
    pf = p.paragraph_format
    pf.space_before = Pt(13 if level == 1 else 8)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    return p


def set_picture_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_image(doc, image_path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Cm(width_cm))
    set_picture_alt(shape, caption, caption)
    p.paragraph_format.space_after = Pt(2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(caption)
    set_run_font(cap_run, 9, italic=True, color=(80, 80, 80))
    cap.paragraph_format.space_after = Pt(8)
    return shape


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.62)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(item)
        set_run_font(run, 10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.text = ""
        set_cell_shading(cell, "243447")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Cm(widths[index])
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = para.add_run(header)
        set_run_font(r, 9.5, bold=True, color=(255, 255, 255))
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_tr_pr.append(header_marker)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        shade = "F1F4F7" if row_index % 2 else "FFFFFF"
        for index, value in enumerate(row):
            cell = cells[index]
            cell.text = ""
            set_cell_shading(cell, shade)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Cm(widths[index])
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.05
            r = para.add_run(value)
            set_run_font(r, 9.5)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_post_block(doc, block):
    add_heading(doc, block["heading"], 1)
    add_para(doc, f"Target forum: {block['where']}", bold_prefix="Target forum: ", space_after=2)
    add_para(doc, f"Post format: {block['format']}", bold_prefix="Post format: ", space_after=2)
    add_para(doc, f"Visual: {block['visual']}", bold_prefix="Visual: ", space_after=2)
    add_para(doc, f"Link use: {block['link_note']}", bold_prefix="Link use: ", space_after=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    lead = p.add_run("Direct URL: ")
    set_run_font(lead, 11, bold=True)
    add_hyperlink(p, block["url"], block["url"])
    add_para(doc, f"Before posting: {block['preflight']}", bold_prefix="Before posting: ", space_after=6)
    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(2)
    r = label.add_run("Ready to post")
    set_run_font(r, 11, bold=True)
    for paragraph_text in block["copy"].split("\n\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.15)
        p.paragraph_format.right_indent = Cm(0.15)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.1
        url_start = paragraph_text.find("https://draxlerwheels.com/")
        if url_start >= 0:
            prefix = paragraph_text[:url_start]
            r = p.add_run(prefix)
            set_run_font(r, 10.4, name="Arial")
            add_hyperlink(p, paragraph_text[url_start:], paragraph_text[url_start:])
        else:
            r = p.add_run(paragraph_text)
            set_run_font(r, 10.4, name="Arial")
    add_para(doc, f"First author reply: {block['reply']}", bold_prefix="First author reply: ", space_after=9)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.75)
section.bottom_margin = Cm(1.65)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.page_width = Cm(21.59)
section.page_height = Cm(27.94)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:cs"), "Arial")
styles["Normal"].font.size = Pt(11)
for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    style.font.color.rgb = RGBColor(0, 0, 0)

# Cover
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_before = Pt(40)
title.paragraph_format.space_after = Pt(12)
set_run_font(title.add_run("Draxler Forum Content Plan"), 28, bold=True, color=(0, 0, 0))
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.paragraph_format.space_after = Pt(18)
set_run_font(subtitle.add_run("10 disclosed community post drafts and a practical publishing sequence"), 14, color=(0, 0, 0))
add_para(doc, "Plan for English-language automotive communities. The copy uses a natural forum voice without hiding the commercial interest: the poster always discloses their link to Draxler and never presents a 3D visual as a photo of a completed customer car.", space_after=10)
add_para(doc, "Core idea: give the audience a useful, specific reason to discuss fitment, proportion, brake visibility, finish or function first, then provide a clear route to the catalog. This mirrors the strongest build and fitment threads among BMW, Mercedes, Porsche and Audi owners.", space_after=18)
add_image(doc, ASSETS / "g-wagon.jpg", 15.5, "Draxler site visual  Mercedes AMG G 63 project")
prepared = doc.add_paragraph()
prepared.paragraph_format.space_before = Pt(8)
prepared.paragraph_format.space_after = Pt(0)
set_run_font(prepared.add_run("Prepared 20 October 2026"), 10, color=(80, 80, 80))

doc.add_page_break()

# Strategy and research
add_heading(doc, "How these posts should behave", 1)
add_para(doc, "Forum voice matters more than brand voice here. Strong discussions do not sell ‘perfect wheels’ in the first sentence. The poster frames one real question, names the limitations and leaves room for disagreement. That is why most drafts below open with full disclosure and a 3D visual study, while technical numbers are offered only after the specific vehicle has been checked.", space_after=7)
add_bullets(doc, [
    "No stealth seeding. Put ‘I work with Draxler’ or ‘I’m with the Draxler team’ in the first sentence.",
    "Label every 3D image as a 3D configuration or visual study. Do not call it an owner build, delivered car or customer car.",
    "Do not state ET, width, weight, brake clearance, load rating or lead time until it is confirmed for the individual car. Ask for the car data instead and return with a verified specification.",
    "Use one CTA per post. Direct Draxler links are mandatory in this plan, but keep only one link and comply with the forum’s commercial-link policy.",
    "Post only in an approved vendor section or after moderator approval. If a community does not permit brand representatives, do not evade that rule with a new account or private messages.",
])

add_heading(doc, "What the research shows", 1)
research_rows = [
    ("Bimmerpost build threads", "The strongest responses centre on a specific car, staggered fitment, tyres and a direct answer on rubbing or brake clearance. Visual taste comes after fitment."),
    ("MBWorld Vossen gallery", "A product gallery works within the exact model section and is refreshed with visuals. It suits an approved vendor presence, not a random drive-by thread."),
    ("Rennlist fitment threads", "Trust threshold is high: platform accuracy, separate front and rear details, and clear statements about what is physically installed."),
    ("Audizine dealer threads", "Catalog announcements belong in an approved vendor context, but the thread needs updates, answers and facts rather than a one-time advertisement."),
    ("Reddit owner discussions", "Readers quickly ask about diameter, tyres, clearance, maintenance and price. Unsupported hype earns irony rather than leads."),
]
add_table(doc, ["Observed pattern", "What Draxler should do"], research_rows, [4.4, 11.6])

add_heading(doc, "Draxler source material used", 1)
add_para(doc, "Draxler presents Luxury, Sport and Off Road series and a catalog of DRX models. Its website includes visuals for a Mercedes AMG G 63, Rolls Royce Phantom, Porsche and Ford Raptor. The posts below use those assets only as brand visuals or reference images; they are not proof of a completed customer installation.", space_after=7)
source_table = doc.add_table(rows=2, cols=3)
source_table.autofit = False
source_table.style = "Table Grid"
source_images = [
    (ASSETS / "luxury-phantom.jpg", "Luxury site visual"),
    (ASSETS / "sport-porsche.jpg", "Sport site visual"),
    (ASSETS / "offroad-ford.jpg", "Off Road site visual"),
    (DRX101, "DRX 101 catalog image"),
    (DRX301, "DRX 301 catalog image"),
    (ASSETS / "g-wagon.jpg", "G 63 project visual"),
]
for i, (image_path, caption) in enumerate(source_images):
    cell = source_table.cell(i // 3, i % 3)
    set_cell_border(cell)
    set_cell_margins(cell, 90, 90, 90, 90)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Cm(4.9))
    set_picture_alt(shape, caption, caption)
    cp = cell.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(caption)
    set_run_font(cr, 8.5, italic=True, color=(80, 80, 80))
source_header_tr_pr = source_table.rows[0]._tr.get_or_add_trPr()
source_header_marker = OxmlElement("w:tblHeader")
source_header_marker.set(qn("w:val"), "true")
source_header_tr_pr.append(source_header_marker)
doc.add_paragraph().paragraph_format.space_after = Pt(3)

add_heading(doc, "Publishing rhythm", 1)
add_table(doc, ["Week", "Posts", "Purpose"], [
    ("1", "BMW G80 and Mercedes G 63", "Earn discussion through visual judgement rather than product claims."),
    ("2", "Porsche 992 and Audi RS6", "Show that the team can listen to platform-specific fitment preferences."),
    ("3", "BMW X7 and Mercedes S Class", "Shift to luxury proportion and finish without claiming completed builds."),
    ("4", "Tesla Model S and Ford Raptor", "Test two different practical objections: range ride and off-road function."),
    ("5", "Draxler fitment log and engineering AMA", "Convert useful answers into a durable, disclosed reference thread."),
], [2.0, 5.6, 8.4])

doc.add_page_break()

posts = [
    {
        "heading": "01 BMW M3 G80 3D fitment vote",
        "where": "Bimmerpost G80 G82 Wheel Tire Suspension section. Do not start in general discussion; request vendor permission first or publish from an account the community already recognizes.",
        "format": "3D visual study plus one direct question. Around 110 words, with no price and no fitment promise.",
        "visual": "A fresh Draxler 3D configurator screenshot of a G80. Add a clear 3D configuration label on the image. Use a DRX Sport wheel close-up as the second image.",
        "link_note": "Use the direct catalog link in the copy. Keep it to one link and confirm the section permits commercial links before posting.",
        "url": "https://draxlerwheels.com/catalog",
        "preflight": "Do not call the concept a completed car. Before replying, have the VIN year, brake package, ride height and customer tyre preference ready.",
        "copy": "Full disclosure: I’m with Draxler. We put together this G80 study to settle the visual direction before we lock any geometry. This is a 3D configuration, not a photo of a completed customer car.\n\nThe question is deliberately simple: does the brushed face with the darker inner barrel keep the car OEM+ enough, or would you go full satin dark here? I like that the spokes leave the brakes visible without making the front end look busy, but side renders can lie.\n\nWe have not published an offset or tyre package because those will be checked against the actual car, brakes and ride height. I’d rather get roasted for the finish now than promise a fitment from a screen. Would you spec this as a 19 20 staggered setup or keep it square?\n\nDraxler catalog: https://draxlerwheels.com/catalog",
        "reply": "Good point on the sidewall. We will not choose wheel diameter until we have the tyre and suspension plan. The attached image is only the design decision stage; I’ll post confirmed width and ET if the forum allows a build follow-up.",
    },
    {
        "heading": "02 BMW X7 G07 luxury SUV proportion study",
        "where": "Bimmerpost X7 G07 forum, wheel and tyre topic or an existing X7 modification thread. Do not create a standalone brand post in the general BMW subreddit without approval.",
        "format": "3D finish comparison. Around 100 words, built around one choice between two finishes.",
        "visual": "A G07 3D screenshot with two finish variants side by side. Use the Draxler DRX 101 catalog image as the second slide.",
        "link_note": "Include the direct catalog link in the copy. Do not write ‘perfect fitment’ until the vehicle has been checked.",
        "url": "https://draxlerwheels.com/catalog/vip/drx-101",
        "preflight": "Check which factory dimensions and load requirements apply to the specific X7. Do not state 22 or 23 inches as a fixed recommendation until confirmed.",
        "copy": "I’m with Draxler, and this is a 3D visual study rather than a delivered X7. We are trying to keep the wheel large enough for the body without making the car look like it is sitting on a concept-car set.\n\nSame design, two directions: brushed clear face with a dark barrel, or a quieter satin graphite. On a big SUV I usually prefer some contrast so the wheel does not disappear, but the polished route can get loud fast.\n\nFitment will be engineered from the real vehicle and tyre plan, so this is not an offset recommendation. Which one reads more factory-plus to you, and would you rather see a dense luxury pattern or something more open around the brakes?\n\nDraxler catalog: https://draxlerwheels.com/catalog/vip/drx-101",
        "reply": "That is exactly the feedback we needed. We are keeping the visual discussion separate from the engineering sheet; no dimensions will be advertised until the car, brakes and load requirements are in front of us.",
    },
    {
        "heading": "03 Mercedes AMG G 63 project reveal",
        "where": "MBWorld G Class W463 section or a moderator-approved showcase thread. This fits the gallery-compilation format, but only with visible vendor disclosure.",
        "format": "Project visual reveal. Around 105 words, with less technical copy and a more specific question about proportion.",
        "visual": "Draxler Mercedes AMG G 63 site image. The post must caption it ‘Draxler project visual, not a customer delivery photo.’ Add a same-angle 3D configurator screenshot when available.",
        "link_note": "Include the direct project URL in the copy. Do not add WhatsApp, a lead form or multiple links.",
        "url": "https://draxlerwheels.com/catalog/projects",
        "preflight": "Prepare an honest answer about whether the visual is a render or photography. Do not claim a specific diameter, PCD or brake clearance from the image.",
        "copy": "Full disclosure: I’m part of the Draxler team. Posting this as a G 63 visual study because the wheel treatment is doing most of the work here. This is a Draxler project visual, not a customer delivery photo.\n\nWe kept the finish dark and the spoke geometry fairly open so the wheel reads as one solid form instead of a black circle. The thing I’m still unsure about is the stance: does the wheel look properly planted here, or does the G need more sidewall and less visual diameter to stay believable?\n\nNo fitment claim attached to this image. If there is interest, I can follow up with a real spec sheet once we are working from a specific G 63 and brake package.\n\nProject reference: https://draxlerwheels.com/catalog/projects",
        "reply": "Agreed on the sidewall. We would treat that as an actual vehicle decision, not something to copy from the render. The point of this post is to hear whether the design direction belongs on a G before geometry is frozen.",
    },
    {
        "heading": "04 Mercedes S Class luxury finish study",
        "where": "MBWorld S Class W223 section or an approved luxury build thread. Do not frame it as ‘new wheels for every S Class’; focus on one specific visual study.",
        "format": "3D luxury concept. Around 100 words, with restrained language and no aggressive tone.",
        "visual": "A 3D configurator capture on a W223 or Maybach specification. Use the Draxler Luxury site image as a second design-reference frame and label it site visual.",
        "link_note": "Include one direct catalog link in the copy. Do not send readers straight to a messenger.",
        "url": "https://draxlerwheels.com/catalog",
        "preflight": "Verify whether the selected car is an S Class or Maybach and do not use incorrect badges in the mock-up. Do not promise OEM cap compatibility.",
        "copy": "I’m with Draxler. This is a 3D S Class study, not an installed car, and I wanted to ask the people who actually live with these cars where the line is.\n\nThe temptation is to add more finish contrast and more spoke detail. We went the other way: a clean luxury face, enough negative space to see the brakes, and a finish that stays quiet next to all that chrome. My concern is that the design could be too restrained once it is moving.\n\nWould you keep it satin titanium, or add a polished lip detail? We will only set width, offset and tyre after the actual chassis and brake package are checked. This post is just about proportion and taste.\n\nDraxler catalog: https://draxlerwheels.com/catalog",
        "reply": "The OEM-cap point is useful. We will not assume it; centre-bore and cap compatibility are part of the final specification, not a render detail.",
    },
    {
        "heading": "05 Porsche 992 Carrera GTS brake visibility question",
        "where": "Rennlist 992 aftermarket wheel fitment thread or a vendor-approved 992 forum post. If the rules prohibit commercial links, get moderator approval before posting the required direct link.",
        "format": "3D performance visual. Around 120 words, with a deliberately narrow question about spoke density and brake visibility.",
        "visual": "A Draxler 3D screenshot specifically on a 992. Second image: Porsche visual from the Draxler site, captioned ‘site visual and design reference, not a completed install.’",
        "link_note": "The direct Draxler catalog link is required in the copy. Keep it as the only link and do not add a purchase form.",
        "url": "https://draxlerwheels.com/catalog",
        "preflight": "Do not state 20 21, centre lock, Pirelli or PCCB compatibility without confirmation. On Rennlist, accuracy matters more than a polished story.",
        "copy": "Full disclosure: I’m with Draxler. We mocked this up on a 992 as a 3D study before we build a real application. It is not an installed customer car and there are no fitment numbers hidden in the render.\n\nThe brief was brake visibility without making the wheel look too delicate for the rear haunches. I think this spoke count works from the side, but I’m less sure about it at a three-quarter angle where the car already has a lot going on.\n\nFor those who have lived with aftermarket 992 wheels: would you rather see a more open face, or does this keep enough visual mass? Once we are working from a specific car, tyre and brake package we can publish only the confirmed dimensions.\n\nDraxler catalog: https://draxlerwheels.com/catalog",
        "reply": "Completely fair. We will not use the usual generic 20 21 language as a substitute for an actual spec. The next useful update would be a measured application, not another render.",
    },
    {
        "heading": "06 Audi RS6 C8 finish comparison",
        "where": "Audizine C8 RS6 section, ideally inside an existing wheel and tyre discussion or an approved advertiser thread. Confirm current vendor rules on the day you post.",
        "format": "Two 3D renders and a short request for critique. Around 105 words.",
        "visual": "Side-by-side Draxler 3D captures of a C8 with brushed-silver and satin-dark variants. Caption the image ‘3D visual study; final fitment to be validated.’",
        "link_note": "Include the direct catalog link in the copy, not a purchase form. The post should stay about the design decision.",
        "url": "https://draxlerwheels.com/catalog",
        "preflight": "Prepare a question about winter sets and wheel cleaning. Do not write ‘RS6 owners all prefer’ or make another broad generalisation.",
        "copy": "I’m with Draxler. We are using this C8 RS6 3D study to decide whether the car wants contrast or a complete blackout. The two images are not a delivered build and the final fitment will be checked against the actual car.\n\nSilver makes the spoke design easy to read and gives the car a more OE-plus feel. Satin dark feels more aggressive, but I worry that it loses the wheel completely once the car is rolling.\n\nIf this were yours, which direction would you take? Also interested in the practical answer: for people who daily an RS6, do you regret dense spokes because of cleaning, or is that a non-issue once the finish is right?\n\nDraxler catalog: https://draxlerwheels.com/catalog",
        "reply": "Thanks for bringing up cleaning. It is a better question than a quick like or dislike, and it is exactly why we are asking before locking a finish for a real application.",
    },
    {
        "heading": "07 Tesla Model S Plaid aero visual study",
        "where": "Tesla Motors Club Wheels Tires Suspension section. Do not use Reddit as a direct sales channel without explicit approval for a vendor post.",
        "format": "3D design prompt. Around 100 words, openly acknowledging that it is not a road test.",
        "visual": "A Model S 3D configurator screenshot with the wheel visual. Use two angles so readers can see the face and sidewall, and label it 3D configuration.",
        "link_note": "Include one direct Draxler home-page link in the copy, subject to forum rules.",
        "url": "https://draxlerwheels.com/",
        "preflight": "Do not claim an effect on range, noise, ride or weight: a 3D image does not prove it. Ask the audience for real-world concerns.",
        "copy": "Full disclosure: I work with Draxler. This is a Model S Plaid 3D wheel study, not an owner review and definitely not a range claim.\n\nWe tried to give the side profile a cleaner, more technical look without copying an aero cover. The visual is doing what we wanted, but a screen says nothing about ride, range, noise or durability. Those are the reasons I’m asking before this goes any further.\n\nIf you run aftermarket wheels on a Plaid, what was the first compromise you noticed in real life? And visually, would you keep this face this open, or add more surface area so it sits better with the car’s body side?\n\nDraxler: https://draxlerwheels.com/",
        "reply": "That is the right distinction. We will treat the visual as a design discussion only; any weight or efficiency statement would need measured data from a real setup.",
    },
    {
        "heading": "08 Ford Raptor off road functional design check",
        "where": "Ford Raptor Forum wheel and tyre section or a moderator-approved vendor thread. Select a community where brands may answer technical questions.",
        "format": "A site visual plus a 3D reference. Around 115 words, with the first question about function rather than appearance.",
        "visual": "Draxler Off Road Ford site visual and DRX 301 catalog image. Caption them ‘site visual and wheel design reference.’ Add a 3D capture with the customer tyre profile when possible.",
        "link_note": "Include one direct catalog link in the copy. Do not use ‘beadlock capable’ as a promise for a particular finished configuration without engineering confirmation.",
        "url": "https://draxlerwheels.com/catalog",
        "preflight": "Gather the intended tyre, trail use, load, brake package, valve protection, true beadlock requirement and local legality. Do not publish a load rating from a render.",
        "copy": "I’m with Draxler. This is a design study for our Off Road direction, not a product test or a claim that this exact Raptor setup is ready to order.\n\nThe visual question is easy. The functional question is not: if a wheel looks beadlock-inspired but does not solve a real trail requirement, most of you will spot that immediately. Before we spec a real application, what matters most to you - true beadlock function, load rating, valve protection, easy cleaning, or simply keeping enough sidewall?\n\nThe image is there to show the design language. Any final wheel would still need to be engineered around the actual truck, tyre and use case. I would rather collect the non-glamorous requirements now.\n\nDraxler catalog: https://draxlerwheels.com/catalog",
        "reply": "Appreciate the direct answer. We will not confuse a visual direction with a functional certification. The useful next step is a real requirement list, then a confirmed application.",
    },
    {
        "heading": "09 Draxler disclosed fitment log",
        "where": "One approved vendor section on Bimmerpost, MBWorld or Audizine. This is a long-term reference thread, not a post to broadcast across every community at once.",
        "format": "Opening post for an updateable fitment log. Around 115 words, with no inflated promises or discount CTA.",
        "visual": "One Draxler catalog wheel image and one visual study. Add new photos only when the source, publication permission and caption are all clear.",
        "link_note": "Pin one neutral catalog link in the first message if the moderator permits it. Build traffic through useful updates, not repeated links.",
        "url": "https://draxlerwheels.com/catalog",
        "preflight": "Agree on an update cadence with the moderator. Prepare one template: chassis, ride height, wheel width, ET, tyre, brake notes, measured weight when available and image origin.",
        "copy": "Hi all - I’m with Draxler, and with moderator permission I’m opening this as a fitment reference thread, not a stream of catalog drops.\n\nWhen we have a confirmed application, the update will list the chassis, wheel width, offset, tyre, ride height and any brake or clearance note that has actually been checked. If a post is only a 3D study, it will say so in the first line. If the image is a site visual, it will say that too.\n\nWe will not fill the thread with prices, artificial urgency or copied claims. If you have a platform that is poorly served by current wheel options, post the car, brake and tyre basics. Useful edge cases are more valuable to us than a generic compliment.\n\nDraxler catalog: https://draxlerwheels.com/catalog",
        "reply": "That is the standard we are aiming for. The next update should earn its place with real dimensions and an honest caption, otherwise it does not go into the log.",
    },
    {
        "heading": "10 Forged wheel engineering AMA",
        "where": "Vendor-approved technical section or a scheduled moderator-approved AMA. Do not run an AMA without an engineer or product lead who can answer for engineering and production.",
        "format": "Opening AMA post. Around 115 words. This is the only block where it helps to state plainly which questions are welcome.",
        "visual": "One neutral DRX catalog image and one small diagram or 3D screenshot. Do not attach ten renders; the answers are the main content.",
        "link_note": "Include the direct Draxler home-page link in the copy. Keep it as the only commercial link and confirm permission first.",
        "url": "https://draxlerwheels.com/",
        "preflight": "An engineer or product lead must be in the AMA. Agree in advance which subjects the team can answer honestly, including unknowns and limitations.",
        "copy": "I’m with Draxler. With moderator approval, we would like to run a practical forged-wheel AMA here. This is not a launch announcement and we are not going to pretend every question has a glamorous answer.\n\nAsk about design choices, monoblock versus multi-piece intent, finish maintenance, brake clearance, why an offset that looks good in a render may not work on a real car, or what information we need before modelling a custom set. If we do not have a confirmed answer for a chassis, we will say so rather than turn a guess into a sales line.\n\nWe will bring the relevant engineering person into the replies. The goal is to leave behind a useful technical thread, not to convert every comment into a quote request.\n\nDraxler: https://draxlerwheels.com/",
        "reply": "Great question. We will answer it in the thread with the assumptions stated up front. If the answer depends on actual brake dimensions or ride height, we will spell that out instead of giving a universal number.",
    },
]

for index, block in enumerate(posts):
    if index in {2, 4, 6, 8}:
        doc.add_page_break()
    add_post_block(doc, block)

doc.add_page_break()
add_heading(doc, "Links and reference examples", 1)
add_para(doc, "Use these as format references, not as copy. Do not transfer another brand’s claims, measurements, imagery or wording to Draxler. Forum rules change, so check them on the day you post.", space_after=7)

references = [
    ("Draxler home", "https://draxlerwheels.com/", "Source for brand language, configurator entry point and forging claims presented by the brand."),
    ("Draxler catalog", "https://draxlerwheels.com/catalog", "Source for Luxury, Sport, Off Road series, DRX model codes and project visual references."),
    ("MBWorld Vossen Forged Gallery Compilation Thread", "https://mbworld.org/forums/g55-g63-g65-amg-w463/677056-vossen-forged-wheel-gallery-compilation-thread.html", "Example of a platform specific visual gallery thread."),
    ("Rennlist 992 Aftermarket Wheels Official Fitment Thread", "https://rennlist.com/forums/992/1512473-992-2-aftermarket-wheels-the-official-fitment-thread.html", "Example of fitment-led conversation in a model-specific community."),
    ("Audizine Vossen Dealer Thread", "https://audizine.com/threads/vossen-wheels-cv7-cv5-cv4-cv3-cv2-cv1-dual-concave-end-of-summer-sale.551779/", "Example of a clearly commercial vendor format, useful only where approved."),
    ("Bimmerpost G20 330i Wheel Setup", "https://g20.bimmerpost.com/forums/showthread.php?t=1840361", "Example of discussion turning to real tyre, wheel size and rubbing questions."),
]
for label, url, note in references:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_hyperlink(p, label, url)
    p.add_run("  -  ")
    p.add_run(note)
    for run in p.runs:
        if run.text != label:
            set_run_font(run, 10.3)

add_heading(doc, "Final posting checklist", 1)
add_bullets(doc, [
    "A visible disclosure of brand affiliation appears in the first sentence.",
    "Each 3D screenshot is labelled as a visual study, and each site image is labelled as a site visual.",
    "Every width, ET, tyre, load rating and clearance claim is verified or omitted as a fixed fact.",
    "Live forum rules have been reviewed and required moderator permission obtained.",
    "Each post contains one focused question, one direct Draxler link, no manufactured urgency, no fake reviews and no concealed push to private messages.",
])

doc.core_properties.title = "Draxler Forum Content Plan"
doc.core_properties.subject = "Ten disclosed forum post drafts for Draxler Wheels"
doc.core_properties.author = "Draxler"
doc.core_properties.comments = ""
doc.save(OUT)
print(OUT)
